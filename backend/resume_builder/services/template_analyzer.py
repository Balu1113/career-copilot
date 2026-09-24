from pathlib import Path

import pdfplumber
from docx import Document


def analyze_docx_template(file_path):
    document = Document(file_path)

    template_data = {
        "file_type": "docx",
        "page": {},
        "font": {},
        "headings": [],
        "sections": [],
        "layout": {},
        "paragraphs": {},
    }

    if document.sections:
        section = document.sections[0]

        template_data["page"] = {
            "width_inches": round(section.page_width.inches, 2),
            "height_inches": round(section.page_height.inches, 2),
            "margin_top": round(section.top_margin.inches, 2),
            "margin_bottom": round(section.bottom_margin.inches, 2),
            "margin_left": round(section.left_margin.inches, 2),
            "margin_right": round(section.right_margin.inches, 2),
        }

    fonts = []

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                fonts.append(run.font.name)

    if fonts:
        font_counts = {}

        for font in fonts:
            font_counts[font] = font_counts.get(font, 0) + 1

        template_data["font"]["family"] = max(
            font_counts,
            key=font_counts.get,
        )

    heading_styles = {
        "Title",
        "Subtitle",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
    }

    seen_styles = set()

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name

        if style_name not in heading_styles:
            continue

        if style_name in seen_styles:
            continue

        seen_styles.add(style_name)

        heading_data = {
            "style": style_name,
            "alignment": (
                str(paragraph.alignment)
                if paragraph.alignment
                else None
            ),
        }

        for run in paragraph.runs:
            if not run.text.strip():
                continue

            heading_data["font"] = run.font.name

            heading_data["size"] = (
                round(run.font.size.pt, 1)
                if run.font.size
                else None
            )

            heading_data["bold"] = (
                run.bold
                if run.bold is not None
                else False
            )

            break

        template_data["headings"].append(
            heading_data
        )

    spacing_values = []

    for paragraph in document.paragraphs:
        spacing_after = paragraph.paragraph_format.space_after

        if spacing_after:
            spacing_values.append(
                round(spacing_after.pt, 1)
            )

    if spacing_values:
        template_data["paragraphs"][
            "average_spacing_after"
        ] = round(
            sum(spacing_values) / len(spacing_values),
            1,
        )

    bullet_types = []

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name.lower()

        if "list bullet" in style_name:
            bullet_types.append("bullet")

        elif "list number" in style_name:
            bullet_types.append("number")

    template_data["layout"]["list_style"] = (
        bullet_types[0]
        if bullet_types
        else "bullet"
    )

    template_data["layout"]["table_count"] = (
        len(document.tables)
    )

    template_data["layout"][
        "possible_multi_column"
    ] = len(document.tables) > 0

    section_names = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = paragraph.style.name.lower()

        if (
            "heading" in style_name
            or style_name in {"title", "subtitle"}
        ):
            section_names.append(text)

    template_data["sections"] = section_names

    return template_data


def _most_common(values):
    if not values:
        return None

    counts = {}

    for value in values:
        counts[value] = counts.get(value, 0) + 1

    return max(
        counts,
        key=counts.get,
    )


def _normalize_font_name(font_name):
    if not font_name:
        return None

    if "+" in font_name:
        return font_name.split("+", 1)[1]

    return font_name


def _detect_page_size(width, height):
    width_inches = width / 72
    height_inches = height / 72

    if (
        abs(width_inches - 8.27) < 0.15
        and abs(height_inches - 11.69) < 0.15
    ):
        return "A4"

    if (
        abs(width_inches - 8.5) < 0.15
        and abs(height_inches - 11) < 0.15
    ):
        return "LETTER"

    return "CUSTOM"


def _group_words_into_lines(words):
    if not words:
        return []

    sorted_words = sorted(
        words,
        key=lambda word: (
            round(word["top"], 1),
            word["x0"],
        ),
    )

    lines = []

    for word in sorted_words:

        if not lines:
            lines.append(
                {
                    "words": [word],
                    "top": word["top"],
                }
            )
            continue

        current_line = lines[-1]

        if abs(
            word["top"] - current_line["top"]
        ) <= 3:

            current_line["words"].append(word)

        else:

            lines.append(
                {
                    "words": [word],
                    "top": word["top"],
                }
            )

    result = []

    for line in lines:

        words_sorted = sorted(
            line["words"],
            key=lambda word: word["x0"],
        )

        text = " ".join(
            word["text"]
            for word in words_sorted
        )

        sizes = [
            float(word["size"])
            for word in words_sorted
            if word.get("size")
        ]

        fonts = [
            _normalize_font_name(word["fontname"])
            for word in words_sorted
            if word.get("fontname")
        ]

        result.append(
            {
                "text": text,
                "font_size": (
                    round(
                        _most_common(
                            [
                                round(size, 1)
                                for size in sizes
                            ]
                        ),
                        1,
                    )
                    if sizes
                    else None
                ),
                "font": _most_common(fonts),
                "x0": min(
                    word["x0"]
                    for word in words_sorted
                ),
                "x1": max(
                    word["x1"]
                    for word in words_sorted
                ),
                "top": line["top"],
            }
        )

    return result


def _detect_alignment(line, page_width):
    if not line:
        return "left"

    center = (
        line["x0"] + line["x1"]
    ) / 2

    page_center = page_width / 2

    if abs(center - page_center) <= 35:
        return "center"

    if line["x1"] >= page_width - 60:
        return "right"

    return "left"


def _detect_columns(lines, page_width):
    """
    Detect genuine multi-column layouts.

    A few right-positioned lines such as dates should
    not cause a resume to be classified as two-column.
    """

    if not lines:
        return 1

    page_center = page_width / 2

    left_lines = []
    right_lines = []

    for line in lines:

        line_center = (
            line["x0"] + line["x1"]
        ) / 2

        if line_center < page_center - 80:
            left_lines.append(line)

        elif line_center > page_center + 80:
            right_lines.append(line)

    if len(left_lines) < 5 or len(right_lines) < 5:
        return 1

    left_area = sum(
        line["x1"] - line["x0"]
        for line in left_lines
    )

    right_area = sum(
        line["x1"] - line["x0"]
        for line in right_lines
    )

    total_area = left_area + right_area

    if total_area == 0:
        return 1

    left_ratio = left_area / total_area
    right_ratio = right_area / total_area

    if (
        left_ratio >= 0.25
        and right_ratio >= 0.25
    ):
        return 2

    return 1


def analyze_pdf_template(file_path):
    """
    Analyze the visual structure of a PDF resume.

    The extracted information describes the template/design.
    Resume content is not stored as template content.
    """

    template_data = {
        "file_type": "pdf",
        "page": {},
        "typography": {},
        "header": {},
        "section_heading": {},
        "body": {},
        "layout": {},
        "sections": [],
    }

    with pdfplumber.open(file_path) as pdf:

        if not pdf.pages:
            raise ValueError(
                "The PDF contains no pages."
            )

        page = pdf.pages[0]

        width = float(page.width)
        height = float(page.height)

        template_data["page"] = {
            "size": _detect_page_size(
                width,
                height,
            ),
            "width_points": round(width, 2),
            "height_points": round(height, 2),
            "width_inches": round(width / 72, 2),
            "height_inches": round(height / 72, 2),
        }

        chars = page.chars

        if not chars:
            return template_data

        # --------------------------------------------------
        # FONT INFORMATION
        # --------------------------------------------------

        font_sizes = []
        font_names = []

        for char in chars:

            size = char.get("size")

            font_name = _normalize_font_name(
                char.get("fontname")
            )

            if size:
                font_sizes.append(
                    float(size)
                )

            if font_name:
                font_names.append(
                    font_name
                )

        if font_sizes:

            template_data["typography"][
                "body_size"
            ] = round(
                _most_common(font_sizes),
                1,
            )

        if font_names:

            template_data["typography"][
                "font"
            ] = _most_common(font_names)

        # --------------------------------------------------
        # FONT SIZE DISTRIBUTION
        # --------------------------------------------------

        size_counts = {}

        for size in font_sizes:

            rounded_size = round(
                size,
                1,
            )

            size_counts[rounded_size] = (
                size_counts.get(
                    rounded_size,
                    0,
                )
                + 1
            )

        sorted_sizes = sorted(
            size_counts.items(),
            key=lambda item: item[0],
            reverse=True,
        )

        if sorted_sizes:

            template_data["typography"][
                "font_size_distribution"
            ] = [
                {
                    "size": size,
                    "character_count": count,
                }
                for size, count in sorted_sizes[:10]
            ]

        # --------------------------------------------------
        # TEXT LINES
        # --------------------------------------------------

        words = page.extract_words(
            extra_attrs=[
                "fontname",
                "size",
            ]
        )

        lines = _group_words_into_lines(
            words
        )

        if lines:

            # --------------------------------------------------
            # HEADER
            # --------------------------------------------------

            largest_line = max(
                lines,
                key=lambda line: (
                    line["font_size"]
                    if line["font_size"]
                    else 0
                ),
            )

            header_text = (
                largest_line["text"].strip()
            )

            template_data["header"] = {
                "alignment": _detect_alignment(
                    largest_line,
                    width,
                ),
                "font_size": (
                    largest_line["font_size"]
                ),
                "font": largest_line["font"],
            }

            # --------------------------------------------------
            # SECTION HEADINGS
            # --------------------------------------------------

            heading_candidates = []

            body_size = template_data[
                "typography"
            ].get("body_size")

            for line in lines:

                text = line["text"].strip()

                if not text:
                    continue

                size = line["font_size"]

                if size is None:
                    continue

                is_header = (
                    text == header_text
                )

                is_larger = (
                    body_size is None
                    or size > body_size + 1
                )

                is_short = len(text) <= 60

                is_probable_heading = (
                    is_larger
                    and is_short
                    and not is_header
                )

                if is_probable_heading:
                    heading_candidates.append(
                        line
                    )

            if heading_candidates:

                heading_sizes = [
                    line["font_size"]
                    for line in heading_candidates
                    if line["font_size"]
                ]

                heading_fonts = [
                    line["font"]
                    for line in heading_candidates
                    if line["font"]
                ]

                heading_alignments = [
                    _detect_alignment(
                        line,
                        width,
                    )
                    for line in heading_candidates
                ]

                template_data[
                    "section_heading"
                ] = {
                    "font_size": round(
                        _most_common(
                            heading_sizes
                        ),
                        1,
                    ),
                    "font": _most_common(
                        heading_fonts
                    ),
                    "alignment": _most_common(
                        heading_alignments
                    ),
                }

                template_data["sections"] = [
                    line["text"]
                    for line in heading_candidates
                ]

            # --------------------------------------------------
            # BODY
            # --------------------------------------------------

            body_lines = []

            if body_size:

                body_lines = [
                    line
                    for line in lines
                    if (
                        line["font_size"]
                        and abs(
                            line["font_size"]
                            - body_size
                        ) <= 0.5
                    )
                ]

            if body_lines:

                substantial_body_lines = [
                    line
                    for line in body_lines
                    if (
                        line["x1"]
                        - line["x0"]
                    ) > width * 0.25
                ]

                if not substantial_body_lines:
                    substantial_body_lines = (
                        body_lines
                    )

                template_data["body"] = {
                    "font_size": round(
                        _most_common(
                            [
                                line[
                                    "font_size"
                                ]
                                for line in substantial_body_lines
                            ]
                        ),
                        1,
                    ),
                    "alignment": _most_common(
                        [
                            _detect_alignment(
                                line,
                                width,
                            )
                            for line in substantial_body_lines
                        ]
                    ),
                }

            # --------------------------------------------------
            # RIGHT-ALIGNED CONTENT
            # --------------------------------------------------

            right_aligned_lines = [
                line
                for line in lines
                if _detect_alignment(
                    line,
                    width,
                ) == "right"
            ]

            template_data["layout"][
                "right_aligned_content"
            ] = bool(
                right_aligned_lines
            )

        # --------------------------------------------------
        # COLUMN DETECTION
        # --------------------------------------------------

        template_data["layout"][
            "columns"
        ] = _detect_columns(
            lines,
            width,
        )

        template_data["layout"][
            "page_count"
        ] = len(pdf.pages)

    return template_data


def analyze_template(file_path):
    extension = Path(
        file_path
    ).suffix.lower()

    if extension == ".docx":
        return analyze_docx_template(
            file_path
        )

    if extension == ".pdf":
        return analyze_pdf_template(
            file_path
        )

    raise ValueError(
        "Only DOCX and PDF sample resumes are supported."
    )